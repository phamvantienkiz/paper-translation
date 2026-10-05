import fitz
import pytest
from pathlib import Path
from docx import Document
from app.services.parser.pdf_parser import PdfParser
from app.services.parser.docx_parser import DocxParser
from app.services.parser.math_isolator import LaTeXIsolator


def test_pdf_parser_success(tmp_path):
    # 1. Create a dummy scientific PDF with PyMuPDF
    pdf_path = tmp_path / "sample_paper.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Deep Residual Learning for Image Recognition", fontsize=18)
    page.insert_text((50, 120), "1. Introduction", fontsize=14)
    page.insert_text(
        (50, 160),
        "Let $H(x)$ denote an underlying mapping. We let the stacked layers fit $F(x) := H(x) - x$.",
        fontsize=11,
    )
    doc.save(str(pdf_path))
    doc.close()

    parser = PdfParser()
    parsed_doc = parser.parse(Path(pdf_path), "test_session_pdf")

    assert len(parsed_doc.blocks) >= 2

    # Check formula isolation on the blocks
    all_tokens = []
    for b in parsed_doc.blocks:
        masked, tokens = LaTeXIsolator.isolate_formulas(b.content)
        all_tokens.extend(tokens)
    assert len(all_tokens) >= 1


def test_pdf_parser_rejects_empty_or_scanned(tmp_path):
    # Create empty PDF without any text layer
    empty_pdf_path = tmp_path / "empty_scan.pdf"
    doc = fitz.open()
    doc.new_page()  # blank page
    doc.save(str(empty_pdf_path))
    doc.close()

    parser = PdfParser()
    with pytest.raises(ValueError, match="lớp văn bản"):
        parser.parse(Path(empty_pdf_path), "test_session_empty")


def test_docx_parser_success(tmp_path):
    # 1. Create a dummy scientific DOCX
    docx_path = tmp_path / "sample_paper.docx"
    doc = Document()
    doc.add_heading("Attention Is All You Need", level=1)
    doc.add_paragraph("The Transformer is the first transduction model relying entirely on self-attention.")
    doc.add_paragraph("Given query and key matrices, we compute $Attention(Q, K, V) = \\text{softmax}(QK^T / \\sqrt{d_k})V$.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "BLEU Score"
    table.cell(1, 0).text = "Transformer (big)"
    table.cell(1, 1).text = "28.4"

    doc.save(str(docx_path))

    parser = DocxParser()
    parsed_doc = parser.parse(Path(docx_path), "test_session_docx")

    assert len(parsed_doc.blocks) >= 3
    # Check title / heading segment
    heading_segs = [b for b in parsed_doc.blocks if b.segment_type in ("TITLE", "HEADING")]
    assert len(heading_segs) >= 1 or parsed_doc.title is not None

    # Check table segment
    table_segs = [b for b in parsed_doc.blocks if b.segment_type == "TABLE"]
    assert len(table_segs) >= 1
    assert "BLEU Score" in table_segs[0].content
